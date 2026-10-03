# Builds House_79K_Grey_Structure_BOQ_Rev1.xlsx from rev1.json with live formulas.
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

HERE = os.path.dirname(os.path.abspath(__file__))
D = json.load(open(os.path.join(HERE, 'rev1.json')))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, 'House_79K_Grey_Structure_BOQ_Rev1.xlsx')

F = 'Arial'
NAVY = '1F3864'
f_title = Font(name=F, size=14, bold=True, color=NAVY)
f_h = Font(name=F, size=9, bold=True, color='FFFFFF')
f_b = Font(name=F, size=9)
f_bb = Font(name=F, size=9, bold=True)
f_in = Font(name=F, size=9, color='0000FF')          # inputs: blue
f_link = Font(name=F, size=9, color='008000')        # cross-sheet links: green
fill_h = PatternFill('solid', fgColor=NAVY)
fill_in = PatternFill('solid', fgColor='FFFF00')     # editable cells: yellow
fill_stage = PatternFill('solid', fgColor='D9E2F3')
fill_tot = PatternFill('solid', fgColor='C6D9F1')
thin = Side(style='thin', color='A6A6A6')
bd = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical='top')
NUM1 = '#,##0.0;(#,##0.0);"-"'
NUM0 = '#,##0;(#,##0);"-"'
PCT = '0%'
RS = '#,##0;(#,##0);"-"'

wb = Workbook()

def header(ws, row, cols, widths=None):
    for c, h in enumerate(cols, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = f_h; cell.fill = fill_h; cell.border = bd
        cell.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)

# =========================================================== Rates sheet
wr = wb.active; wr.title = 'Rates'
wr['A1'] = 'Rates and wastage - edit the yellow cells (blue text = input)'; wr['A1'].font = f_title
wr['A2'] = 'Rates in PKR per unit, Faisalabad, October 2026. "estimated" = no published 2026 figure found. Wastage per Register K2 (lower-middle of normal range).'
wr['A2'].font = f_b
RH = 4
header(wr, RH, ['Key', 'Material', 'Unit', 'Wastage %', 'Rate (PKR / unit)', 'Buying unit', 'Unit size', 'Rate date', 'Source / note', 'Group'],
       [9, 52, 7, 10, 14, 18, 9, 12, 70, 14])
keys = list(D['rates'].keys())
RATE_ROW = {}
for r, k in enumerate(keys, RH + 1):
    name, unit, bunit, usize, w, rate, date, src, grp = D['rates'][k]
    vals = [k, name, unit, w, rate, bunit, usize, date, src, grp]
    for c, v in enumerate(vals, 1):
        cell = wr.cell(row=r, column=c, value=v); cell.font = f_b; cell.border = bd; cell.alignment = wrap
    for c in (4, 5):
        wr.cell(row=r, column=c).font = f_in; wr.cell(row=r, column=c).fill = fill_in
    wr.cell(row=r, column=4).number_format = PCT
    wr.cell(row=r, column=5).number_format = '#,##0.00'
    RATE_ROW[k] = r
RLAST = RH + len(keys)
wr.cell(row=RLAST + 2, column=1, value='Legend: yellow cells with blue text are inputs; all other numbers in this workbook are formulas.').font = f_bb
wr.cell(row=RLAST + 3, column=1, value='Example: to price cement at Rs 1,600 per bag, type 1600 in the cement row of column E; every sheet updates.').font = f_b
rk = f"Rates!$A${RH+1}:$A${RLAST}"
def rlook(col, keycell):
    return f"IFERROR(INDEX(Rates!${col}${RH+1}:${col}${RLAST},MATCH({keycell},{rk},0)),0)"

# =========================================================== BOQ sheet
wbq = wb.create_sheet('BOQ')
wbq['A1'] = 'House 79/K - Grey structure MATERIAL BOQ, Revision 1 (by construction stage)'; wbq['A1'].font = f_title
wbq['A2'] = ('One row per material of each BOQ line. Net qty = measured quantity; Gross = Net x (1 + wastage); Amount = Gross x Rate. '
             'Wastage and rates come from the Rates sheet. Purchase rounding to buying units is on the Totals sheet.')
wbq['A2'].font = f_b
BH = 4
cols = ['Stage', 'Stage name', 'Item', 'Location', 'Description', 'Dimensions', 'Quantity working', 'Item qty', 'Item unit',
        'Register', 'Member group', 'Earth tag', 'Mat. key', 'Material', 'Net qty', 'Unit', 'Wastage %', 'Gross qty', 'Rate (PKR)',
        'Amount (PKR)', 'Loads (trolleys)']
header(wbq, BH, cols, [6, 16, 7, 18, 46, 30, 30, 10, 6, 10, 13, 8, 8, 30, 11, 6, 8, 11, 10, 13, 9])
STAGE_NAME = dict((s, n) for s, n in D['stages'])
row = BH + 1
first_row = row
EARTH_ROW = None
for it in D['items']:
    mats = list(it['mat'].items())
    tag = ''
    for t in ('EXC', 'BACKFILL', 'EFILL'):
        if it['info'].get(t): tag = t
    if not mats: mats = [(None, None)]
    for j, (k, v) in enumerate(mats):
        r = row
        base = [it['stage'], STAGE_NAME[it['stage']], it['code']]
        for c, val in enumerate(base, 1):
            wbq.cell(row=r, column=c, value=val)
        if j == 0:
            wbq.cell(row=r, column=4, value=it['loc']); wbq.cell(row=r, column=5, value=it['desc'])
            wbq.cell(row=r, column=6, value=it['dims']); wbq.cell(row=r, column=7, value=it['work'])
            wbq.cell(row=r, column=8, value=round(it['qty'], 3)); wbq.cell(row=r, column=9, value=it['unit'])
            wbq.cell(row=r, column=12, value=tag)
        wbq.cell(row=r, column=10, value=it['reg']); wbq.cell(row=r, column=11, value=it['group'])
        if k:
            wbq.cell(row=r, column=13, value=k)
            wbq.cell(row=r, column=14, value=f"=IFERROR(INDEX(Rates!$B${RH+1}:$B${RLAST},MATCH(M{r},{rk},0)),\"\")")
            if it['ref'] == '6.9':     # earth to buy: live earth balance
                EARTH_ROW = r
                wbq.cell(row=r, column=15, value=f'=MAX(0,SUMIF($L${first_row}:$L$2000,"EFILL",$H${first_row}:$H$2000)-(SUMIF($L${first_row}:$L$2000,"EXC",$H${first_row}:$H$2000)-SUMIF($L${first_row}:$L$2000,"BACKFILL",$H${first_row}:$H$2000))*1.25)')
                wbq.cell(row=r, column=15).comment = Comment('Earth to buy = fill required (loose) - surplus excavated earth x 1.25 bulking. Live formula from the EXC / BACKFILL / EFILL rows.', 'estimate')
                wbq.cell(row=r, column=8, value=f'=O{r}')
            else:
                wbq.cell(row=r, column=15, value=round(v, 4))
            wbq.cell(row=r, column=16, value=f"=IFERROR(INDEX(Rates!$C${RH+1}:$C${RLAST},MATCH(M{r},{rk},0)),\"\")")
            wbq.cell(row=r, column=17, value='=' + rlook('D', f'M{r}'))
            wbq.cell(row=r, column=18, value=f'=O{r}*(1+Q{r})')
            wbq.cell(row=r, column=19, value='=' + rlook('E', f'M{r}'))
            wbq.cell(row=r, column=20, value=f'=R{r}*S{r}')
            wbq.cell(row=r, column=21, value=f'=IF(M{r}="BRK",R{r}/2000,IF(OR(M{r}="SNDC",M{r}="SNDM",M{r}="CRS",M{r}="KHG",M{r}="EARTH"),R{r}/100,""))')
        else:
            wbq.cell(row=r, column=20, value=0)
        for c in range(1, 22):
            cell = wbq.cell(row=r, column=c); cell.font = f_b; cell.border = bd
            if c in (4, 5, 6, 7, 14): cell.alignment = wrap
        for c in (8, 15, 18): wbq.cell(row=r, column=c).number_format = NUM1
        wbq.cell(row=r, column=17).number_format = PCT
        wbq.cell(row=r, column=19).number_format = '#,##0.00'
        wbq.cell(row=r, column=20).number_format = RS
        wbq.cell(row=r, column=21).number_format = '0.0'
        for c in (14, 16, 17, 19): wbq.cell(row=r, column=c).font = f_link
        row += 1
last_row = row - 1
wbq.cell(row=row + 1, column=5, value='Total of BOQ lines (gross x rate, before purchase rounding; binding wire is on the Totals sheet)').font = f_bb
wbq.cell(row=row + 1, column=20, value=f'=SUM(T{first_row}:T{last_row})').font = f_bb
wbq.cell(row=row + 1, column=20).number_format = RS
BOQ_TOTAL_CELL = f'BOQ!T{row+1}'
wbq.auto_filter.ref = f'A{BH}:U{last_row}'
rng = lambda col: f"BOQ!${col}${first_row}:${col}${last_row}"

# =========================================================== Totals sheet
wt = wb.create_sheet('Totals')
wt['A1'] = 'Totals by material - purchase quantities rounded to buying units'; wt['A1'].font = f_title
wt['A2'] = 'Net = SUMIF of BOQ; Gross = Net x (1 + wastage); Purchase = Gross rounded UP to the buying unit; Cost = Purchase x Rate.'; wt['A2'].font = f_b
TH = 4
header(wt, TH, ['Key', 'Material', 'Unit', 'Net qty', 'Wastage %', 'Gross qty', 'Buying unit', 'Unit size', 'Purchase qty', 'No. of buying units',
                'Rate (PKR)', 'Cost (PKR)', 'Tractor-trolley loads', 'Group'],
       [9, 52, 7, 12, 9, 12, 18, 9, 12, 11, 11, 14, 11, 14])
TROW = {}
skeys = [s['key'] for s in D['summary']]
for r, k in enumerate(skeys, TH + 1):
    TROW[k] = r
STEEL_GROSS_CELLS = [f'F{TROW[k]}' for k in ('ST2', 'ST3', 'ST4', 'ST6') if k in TROW]
for k in skeys:
    r = TROW[k]
    wt.cell(row=r, column=1, value=k)
    wt.cell(row=r, column=2, value=f"=INDEX(Rates!$B${RH+1}:$B${RLAST},MATCH(A{r},{rk},0))")
    wt.cell(row=r, column=3, value=f"=INDEX(Rates!$C${RH+1}:$C${RLAST},MATCH(A{r},{rk},0))")
    if k == 'BWIRE':
        wt.cell(row=r, column=4, value='=(' + '+'.join(STEEL_GROSS_CELLS) + ')*0.01')
        wt.cell(row=r, column=4).comment = Comment('Binding wire = 10 kg per ton of steel (gross), Register E4.', 'estimate')
    else:
        wt.cell(row=r, column=4, value=f'=SUMIF({rng("M")},A{r},{rng("O")})')
    wt.cell(row=r, column=5, value='=' + rlook('D', f'A{r}'))
    wt.cell(row=r, column=6, value=f'=D{r}*(1+E{r})')
    wt.cell(row=r, column=7, value=f"=INDEX(Rates!$F${RH+1}:$F${RLAST},MATCH(A{r},{rk},0))")
    wt.cell(row=r, column=8, value=f"=INDEX(Rates!$G${RH+1}:$G${RLAST},MATCH(A{r},{rk},0))")
    wt.cell(row=r, column=9, value=f'=ROUNDUP(F{r}/H{r},0)*H{r}')
    wt.cell(row=r, column=10, value=f'=I{r}/H{r}')
    wt.cell(row=r, column=11, value='=' + rlook('E', f'A{r}'))
    wt.cell(row=r, column=12, value=f'=I{r}*K{r}')
    wt.cell(row=r, column=13, value=f'=IF(A{r}="BRK",ROUNDUP(I{r}/2000,0),IF(OR(A{r}="SNDC",A{r}="SNDM",A{r}="CRS",A{r}="KHG",A{r}="EARTH"),ROUNDUP(I{r}/100,0),""))')
    wt.cell(row=r, column=14, value=f"=INDEX(Rates!$J${RH+1}:$J${RLAST},MATCH(A{r},{rk},0))")
    for c in range(1, 15):
        cell = wt.cell(row=r, column=c); cell.font = f_b; cell.border = bd
    for c in (2, 3, 5, 7, 8, 11, 14): wt.cell(row=r, column=c).font = f_link
    for c in (4, 6, 9): wt.cell(row=r, column=c).number_format = NUM1
    wt.cell(row=r, column=5).number_format = PCT; wt.cell(row=r, column=10).number_format = '#,##0.0'
    wt.cell(row=r, column=11).number_format = '#,##0.00'; wt.cell(row=r, column=12).number_format = RS
TLAST = TH + len(skeys)
r = TLAST + 2
wt.cell(row=r, column=2, value='TOTAL MATERIAL COST (purchase quantities)').font = f_bb
wt.cell(row=r, column=12, value=f'=SUM(L{TH+1}:L{TLAST})'); TOTAL_CELL = f'Totals!L{r}'
wt.cell(row=r + 1, column=2, value='Covered area (sft) - GF 1,898 + FF 1,844 + mumty 362').font = f_b
wt.cell(row=r + 1, column=12, value=D['covered']); wt.cell(row=r + 1, column=12).font = f_in; wt.cell(row=r + 1, column=12).fill = fill_in
wt.cell(row=r + 2, column=2, value='Cost per sft of covered area (PKR)').font = f_bb
wt.cell(row=r + 2, column=12, value=f'=L{r}/L{r+1}')
wt.cell(row=r + 3, column=2, value='Total of BOQ lines before purchase rounding (incl. binding wire gross)').font = f_b
wt.cell(row=r + 3, column=12, value=f"={BOQ_TOTAL_CELL}+F{TROW['BWIRE']}*K{TROW['BWIRE']}")
wt.cell(row=r + 4, column=2, value='Rounding to buying units').font = f_b
wt.cell(row=r + 4, column=12, value=f'=L{r}-L{r+3}')
for rr in range(r, r + 5):
    wt.cell(row=rr, column=12).number_format = RS; wt.cell(row=rr, column=12).border = bd
    if rr == r or rr == r + 2:
        wt.cell(row=rr, column=12).font = f_bb; wt.cell(row=rr, column=12).fill = fill_tot
TOT_ROW = r
# category subtotals
r = TOT_ROW + 7
wt.cell(row=r, column=2, value='Cost by group').font = f_bb
groups = []
for s in D['summary']:
    if s['group'] not in groups: groups.append(s['group'])
for g in groups:
    r += 1
    wt.cell(row=r, column=2, value=g).font = f_b
    wt.cell(row=r, column=12, value=f'=SUMIF(N{TH+1}:N{TLAST},B{r},L{TH+1}:L{TLAST})').number_format = RS

# =========================================================== Stages sheet
wsg = wb.create_sheet('Stages')
wsg['A1'] = 'Stage-wise cost and quantities to buy before each stage (gross = with wastage)'; wsg['A1'].font = f_title
SH = 3
skeys2 = ['CEM', 'BRK', 'ST2', 'ST3', 'ST4', 'ST6', 'SNDC', 'SNDM', 'CRS', 'KHG', 'EARTH']
labels = ['Cement (bags)', 'Bricks (nos)', 'Steel #2 (kg)', 'Steel #3 (kg)', 'Steel #4 (kg)', 'Steel #6 (kg)', 'Chenab sand (cft)',
          'Ravi sand (cft)', 'Crush (cft)', 'Khangar (cft)', 'Earth (cft)']
header(wsg, SH, ['Stage', 'Stage name', 'Material cost (PKR, before rounding)'] + labels, [7, 40, 16] + [11]*len(labels))
for i, (s, n) in enumerate(D['stages']):
    r = SH + 1 + i
    wsg.cell(row=r, column=1, value=s); wsg.cell(row=r, column=2, value=n)
    wsg.cell(row=r, column=3, value=f'=SUMIF({rng("A")},A{r},{rng("T")})').number_format = RS
    for j, k in enumerate(skeys2):
        c = wsg.cell(row=r, column=4 + j, value=f'=SUMIFS({rng("R")},{rng("A")},$A{r},{rng("M")},"{k}")')
        c.number_format = NUM0
    for c in range(1, 4 + len(skeys2)):
        wsg.cell(row=r, column=c).font = f_b; wsg.cell(row=r, column=c).border = bd
SL = SH + len(D['stages'])
r = SL + 1
wsg.cell(row=r, column=2, value='Binding wire (bought with the steel)').font = f_b
wsg.cell(row=r, column=3, value=f"=Totals!F{TROW['BWIRE']}*Totals!K{TROW['BWIRE']}").number_format = RS
r += 1
wsg.cell(row=r, column=2, value='TOTAL').font = f_bb
wsg.cell(row=r, column=3, value=f'=SUM(C{SH+1}:C{SL+1})').number_format = RS
for j in range(len(skeys2)):
    col = get_column_letter(4 + j)
    wsg.cell(row=r, column=4 + j, value=f'=SUM({col}{SH+1}:{col}{SL})').number_format = NUM0
for c in range(1, 4 + len(skeys2)):
    wsg.cell(row=r, column=c).font = f_bb; wsg.cell(row=r, column=c).fill = fill_tot; wsg.cell(row=r, column=c).border = bd

# =========================================================== Steel sheet
wst = wb.create_sheet('Steel')
wst['A1'] = 'Steel schedule by bar size and member group (net kg, before 4% wastage)'; wst['A1'].font = f_title
STH = 3
header(wst, STH, ['Member group', '#2 (6 mm)', '#3 (10 mm)', '#4 (12 mm)', '#6 (20 mm)', 'Total (kg)'], [22, 12, 12, 12, 12, 12])
for i, (g, _) in enumerate(D['steel_rows']):
    r = STH + 1 + i
    wst.cell(row=r, column=1, value=g)
    for j, k in enumerate(['ST2', 'ST3', 'ST4', 'ST6']):
        wst.cell(row=r, column=2 + j, value=f'=SUMIFS({rng("O")},{rng("K")},$A{r},{rng("M")},"{k}")').number_format = NUM1
    wst.cell(row=r, column=6, value=f'=SUM(B{r}:E{r})').number_format = NUM1
    for c in range(1, 7): wst.cell(row=r, column=c).font = f_b; wst.cell(row=r, column=c).border = bd
SLAST = STH + len(D['steel_rows'])
r = SLAST + 1
wst.cell(row=r, column=1, value='TOTAL net').font = f_bb
for c in range(2, 7):
    col = get_column_letter(c)
    wst.cell(row=r, column=c, value=f'=SUM({col}{STH+1}:{col}{SLAST})').number_format = NUM1
    wst.cell(row=r, column=c).font = f_bb; wst.cell(row=r, column=c).fill = fill_tot
r += 1
wst.cell(row=r, column=1, value='Purchase (Totals sheet)').font = f_bb
for j, k in enumerate(['ST2', 'ST3', 'ST4', 'ST6']):
    wst.cell(row=r, column=2 + j, value=f'=Totals!I{TROW[k]}').number_format = NUM0
wst.cell(row=r, column=6, value=f'=SUM(B{r}:E{r})').number_format = NUM0
r += 1
wst.cell(row=r, column=1, value='Binding wire purchase (kg)').font = f_b
wst.cell(row=r, column=6, value=f"=Totals!I{TROW['BWIRE']}").number_format = NUM0

# =========================================================== Levels sheet
wl = wb.create_sheet('Levels')
wl['A1'] = 'Level table: foundation bottom to finished floor (road crown = 0)'; wl['A1'].font = f_title
header(wl, 3, ['Level (inches)', 'Level (ft-in)', 'What is at this level', 'Register'], [12, 14, 80, 12])
def ftin(x):
    s = '-' if x < 0 else '+'
    a = abs(x); f = int(a // 12); i = a - f*12
    return '%s%d\'-%s"' % (s, f, ('%g' % i))
for i, (lv, txt, reg) in enumerate(sorted(D['levels'], key=lambda x: x[0])):
    r = 4 + i
    for c, v in enumerate([lv, ftin(lv), txt, reg], 1):
        cell = wl.cell(row=r, column=c, value=v); cell.font = f_b; cell.border = bd
r = 5 + len(D['levels'])
wl.cell(row=r, column=1, value='Check B5: top of DPC = finished floor +2\'-9"; highest outside level = porch +1\'-3"; clearance 1\'-6" (minimum 9") - PASS').font = f_bb

# =========================================================== Summary sheet (first)
ws = wb.create_sheet('Summary', 0)
ws['A1'] = 'House No. 79/K, WAPDA City Faisalabad - Grey Structure MATERIAL Estimate, Rev 1'; ws['A1'].font = f_title
ws['A2'] = 'Material only (no labour). Quantities from the structural and working drawings and the owner\'s Final Decisions Register. Rates Oct 2026.'
ws['A2'].font = f_b
ws.column_dimensions['A'].width = 46; ws.column_dimensions['B'].width = 18; ws.column_dimensions['C'].width = 14; ws.column_dimensions['D'].width = 18
rows = [('Total material cost (PKR)', f'={TOTAL_CELL}', RS), ('Covered area (sft)', f'=Totals!L{TOT_ROW+1}', NUM0),
        ('Cost per sft of covered area (PKR)', f'=Totals!L{TOT_ROW+2}', RS)]
for i, (a, f_, nf) in enumerate(rows):
    r = 4 + i
    ws.cell(row=r, column=1, value=a).font = f_bb
    c = ws.cell(row=r, column=2, value=f_); c.number_format = nf; c.font = f_bb; c.fill = fill_tot; c.border = bd
r = 8
for c, h in enumerate(['Main material', 'Purchase qty', 'Unit', 'Cost (PKR)'], 1):
    cell = ws.cell(row=r, column=c, value=h); cell.font = f_h; cell.fill = fill_h; cell.border = bd
main = [('CEM', 'Cement OPC 50 kg'), ('BRK', 'Bricks first class'), ('ST3', 'Steel #3 (10 mm)'), ('ST4', 'Steel #4 (12 mm)'),
        ('ST6', 'Steel #6 (20 mm)'), ('ST2', 'Steel #2 (6 mm)'), ('BWIRE', 'Binding wire'), ('CRS', 'Crush ¾" Sargodha'),
        ('SNDC', 'Chenab sand'), ('SNDM', 'Ravi sand'), ('KHG', 'Khangar (stone ballast)'), ('EARTH', 'Earth to buy'),
        ('BIT', 'Hot bitumen'), ('MEMB', 'Roof membrane 4 mm'), ('TILE', 'Bhatta roof tiles')]
for i, (k, lab) in enumerate(main):
    rr = r + 1 + i
    ws.cell(row=rr, column=1, value=lab)
    ws.cell(row=rr, column=2, value=f'=Totals!I{TROW[k]}').number_format = NUM0
    ws.cell(row=rr, column=3, value=f'=Totals!C{TROW[k]}')
    ws.cell(row=rr, column=4, value=f'=Totals!L{TROW[k]}').number_format = RS
    for c in range(1, 5): ws.cell(row=rr, column=c).font = f_b; ws.cell(row=rr, column=c).border = bd
rr = r + 2 + len(main)
ws.cell(row=rr, column=1, value='Sheets: BOQ (lines, live formulas) | Totals (purchase) | Rates (edit yellow cells) | Stages | Steel | Levels').font = f_b
ws.cell(row=rr + 1, column=1, value='Inputs: yellow cells with blue text on the Rates sheet (rate, wastage) and the covered area on Totals.').font = f_b

wb.save(OUT)
print('saved', OUT, 'BOQ rows', last_row - first_row + 1, 'earth row', EARTH_ROW, 'total cell', TOTAL_CELL)
